"""格式服务：把知识库文档的正文（Markdown）导出为可下载的文档文件。

- docx：python-docx 重建 Word（标题/段落/列表/表格/加粗）
- pdf ：PyMuPDF 生成 PDF（内置 china-s 中文字体，无需外部字体文件）
- md  ：原样输出 Markdown
- txt ：去标记的纯文本
下载文件名 = 原文件名 + (V{版本号})，扩展名随目标格式。

正文以 Markdown 作为唯一交换格式：前端富文本编辑器 → Markdown → 本文档。
"""
from __future__ import annotations

import io
import re
from pathlib import Path

# --------------------------------------------------------------- 块解析
Block = tuple  # ("h", level, text) | ("p", text) | ("ul", items) | ("ol", items)
#              | ("quote", text) | ("code", text) | ("table", rows) | ("hr",)


def parse_blocks(md: str) -> list[Block]:
    """把 Markdown 解析成块序列（行式解析，够用且无第三方依赖）。"""
    lines = (md or "").replace("\r\n", "\n").replace("\r", "\n").split("\n")
    blocks: list[Block] = []
    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        s = line.strip()

        if not s:
            i += 1
            continue

        # 代码块
        if s.startswith("```"):
            i += 1
            buf = []
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1  # 跳过结束 ```
            blocks.append(("code", "\n".join(buf)))
            continue

        # 分隔线
        if re.fullmatch(r"([-*_])\1{2,}", s.replace(" ", "")):
            blocks.append(("hr",))
            i += 1
            continue

        # 标题
        m = re.match(r"^(#{1,6})\s+(.*)$", s)
        if m:
            blocks.append(("h", len(m.group(1)), m.group(2).strip()))
            i += 1
            continue

        # 表格：以 | 开头且下一行是分隔行
        if s.startswith("|") and i + 1 < n and re.match(r"^\s*\|?[\s:|-]+\|?\s*$", lines[i + 1]):
            rows: list[list[str]] = []
            while i < n and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not re.fullmatch(r"[\s:|-]*", "".join(cells)):
                    rows.append(cells)
                i += 1
            if rows:
                blocks.append(("table", rows))
            continue

        # 引用
        if s.startswith(">"):
            buf = []
            while i < n and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip().lstrip(">").strip())
                i += 1
            blocks.append(("quote", " ".join(buf)))
            continue

        # 无序列表
        if re.match(r"^\s*[-*+]\s+", line):
            items = []
            while i < n and re.match(r"^\s*[-*+]\s+", lines[i]):
                items.append(re.sub(r"^\s*[-*+]\s+", "", lines[i]).strip())
                i += 1
            blocks.append(("ul", items))
            continue

        # 有序列表
        if re.match(r"^\s*\d+[.)]\s+", line):
            items = []
            while i < n and re.match(r"^\s*\d+[.)]\s+", lines[i]):
                items.append(re.sub(r"^\s*\d+[.)]\s+", "", lines[i]).strip())
                i += 1
            blocks.append(("ol", items))
            continue

        # 普通段落（连续非空行合并）
        buf = [s]
        i += 1
        while i < n:
            nxt = lines[i].strip()
            if (not nxt or nxt.startswith(("#", "|", ">", "```"))
                    or re.match(r"^\s*[-*+]\s+", lines[i])
                    or re.match(r"^\s*\d+[.)]\s+", lines[i])
                    or re.fullmatch(r"([-*_])\1{2,}", nxt.replace(" ", ""))):
                break
            buf.append(nxt)
            i += 1
        blocks.append(("p", " ".join(buf)))
    return blocks


_INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def strip_inline(text: str) -> str:
    """去掉行内标记，得到纯文本。"""
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    text = re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    return text


def _inline_html(text: str) -> str:
    from html import escape

    def repl(m: re.Match) -> str:
        tok = m.group(0)
        if tok.startswith("**"):
            return f"<strong>{escape(tok[2:-2])}</strong>"
        if tok.startswith("`"):
            return f"<code>{escape(tok[1:-1])}</code>"
        if tok.startswith("*"):
            return f"<em>{escape(tok[1:-1])}</em>"
        lm = re.match(r"\[([^\]]+)\]\(([^)]+)\)", tok)
        if lm:
            return f'<a href="{escape(lm.group(2))}">{escape(lm.group(1))}</a>'
        return escape(tok)

    out, last = [], 0
    for m in _INLINE_RE.finditer(text):
        out.append(escape(text[last:m.start()]))
        out.append(repl(m))
        last = m.end()
    out.append(escape(text[last:]))
    return "".join(out)


# --------------------------------------------------------------- Markdown（原样）
def md_bytes(md: str) -> bytes:
    return (md or "").encode("utf-8")


# --------------------------------------------------------------- 纯文本
def txt_bytes(md: str) -> bytes:
    out_lines: list[str] = []
    for b in parse_blocks(md):
        kind = b[0]
        if kind == "h":
            out_lines.append(strip_inline(b[2]))
            out_lines.append("")
        elif kind == "p":
            out_lines.append(strip_inline(b[1]))
            out_lines.append("")
        elif kind in ("ul", "ol"):
            for it in b[1]:
                out_lines.append("· " + strip_inline(it))
            out_lines.append("")
        elif kind == "quote":
            out_lines.append("> " + strip_inline(b[1]))
            out_lines.append("")
        elif kind == "code":
            out_lines.append(b[1])
            out_lines.append("")
        elif kind == "table":
            for row in b[1]:
                out_lines.append(" | ".join(strip_inline(c) for c in row))
            out_lines.append("")
        elif kind == "hr":
            out_lines.append("----------")
            out_lines.append("")
    return "\n".join(out_lines).strip().encode("utf-8")


# --------------------------------------------------------------- HTML
def html_bytes(md: str, title: str = "文档") -> bytes:
    from html import escape

    parts: list[str] = []
    for b in parse_blocks(md):
        kind = b[0]
        if kind == "h":
            lv = min(b[1], 6)
            parts.append(f"<h{lv}>{_inline_html(b[2])}</h{lv}>")
        elif kind == "p":
            parts.append(f"<p>{_inline_html(b[1])}</p>")
        elif kind == "ul":
            parts.append("<ul>" + "".join(f"<li>{_inline_html(x)}</li>" for x in b[1]) + "</ul>")
        elif kind == "ol":
            parts.append("<ol>" + "".join(f"<li>{_inline_html(x)}</li>" for x in b[1]) + "</ol>")
        elif kind == "quote":
            parts.append(f"<blockquote>{_inline_html(b[1])}</blockquote>")
        elif kind == "code":
            parts.append(f"<pre><code>{escape(b[1])}</code></pre>")
        elif kind == "table":
            rows = b[1]
            head = rows[0]
            body = rows[1:]
            th = "".join(f"<th>{_inline_html(c)}</th>" for c in head)
            trs = "".join(
                "<tr>" + "".join(f"<td>{_inline_html(c)}</td>" for c in r) + "</tr>"
                for r in body
            )
            parts.append(f"<table><thead><tr>{th}</tr></thead><tbody>{trs}</tbody></table>")
        elif kind == "hr":
            parts.append("<hr/>")

    doc = (
        "<!doctype html><html lang=\"zh-CN\"><head><meta charset=\"utf-8\">"
        f"<title>{escape(title)}</title><style>"
        "body{max-width:820px;margin:40px auto;padding:0 20px;font-family:"
        "-apple-system,'Segoe UI','Microsoft YaHei',sans-serif;line-height:1.85;color:#222}"
        "h1{font-size:26px}h2{font-size:21px}h3{font-size:17px;margin-top:1.4em}"
        "table{border-collapse:collapse;margin:14px 0;width:100%}"
        "th,td{border:1px solid #ccc;padding:8px 12px;text-align:left}"
        "th{background:#f3f4f8}code{background:#f0f0f4;padding:1px 5px;border-radius:4px}"
        "pre{background:#f6f6fa;padding:14px;border-radius:8px;overflow:auto}"
        "blockquote{border-left:4px solid #c9c9d6;margin:12px 0;padding:4px 16px;color:#555}"
        "</style></head><body>" + "".join(parts) + "</body></html>"
    )
    return doc.encode("utf-8")


# --------------------------------------------------------------- Word
def docx_bytes(md: str) -> bytes:
    from docx import Document as DocxDocument
    from docx.shared import Pt

    doc = DocxDocument()
    doc.styles["Normal"].font.size = Pt(11)

    def add_inline(paragraph, text: str) -> None:
        for tok in _INLINE_RE.split(text):
            if not tok:
                continue
            if tok.startswith("**") and tok.endswith("**"):
                paragraph.add_run(tok[2:-2]).bold = True
            elif tok.startswith("`") and tok.endswith("`"):
                paragraph.add_run(tok[1:-1]).font.name = "Consolas"
            elif tok.startswith("*") and tok.endswith("*"):
                paragraph.add_run(tok[1:-1]).italic = True
            else:
                lm = re.match(r"\[([^\]]+)\]\(([^)]+)\)", tok)
                paragraph.add_run(lm.group(1) if lm else strip_inline(tok))

    for b in parse_blocks(md):
        kind = b[0]
        if kind == "h":
            doc.add_heading(strip_inline(b[2]), level=min(b[1], 4))
        elif kind == "p":
            add_inline(doc.add_paragraph(), b[1])
        elif kind == "ul":
            for it in b[1]:
                add_inline(doc.add_paragraph(style="List Bullet"), it)
        elif kind == "ol":
            for it in b[1]:
                add_inline(doc.add_paragraph(style="List Number"), it)
        elif kind == "quote":
            add_inline(doc.add_paragraph(style="Intense Quote"), b[1])
        elif kind == "code":
            p = doc.add_paragraph()
            run = p.add_run(b[1])
            run.font.name = "Consolas"
        elif kind == "table":
            rows = b[1]
            cols = max(len(r) for r in rows)
            t = doc.add_table(rows=0, cols=cols)
            t.style = "Table Grid"
            for r in rows:
                cells = t.add_row().cells
                for i in range(cols):
                    cells[i].text = strip_inline(r[i]) if i < len(r) else ""
        elif kind == "hr":
            doc.add_paragraph("─" * 30)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------- PDF
_MARGIN = 72.0
_PAGE_W, _PAGE_H = 595.0, 842.0  # A4 pt
_BOTTOM = _PAGE_H - _MARGIN


class _Pdf:
    def __init__(self) -> None:
        import pymupdf

        self.fitz = pymupdf
        self.doc = pymupdf.open()
        self.page = self.doc.new_page(width=_PAGE_W, height=_PAGE_H)
        self.y = _MARGIN

    def new_page(self) -> None:
        self.page = self.doc.new_page(width=_PAGE_W, height=_PAGE_H)
        self.y = _MARGIN

    def text(self, content: str, size: float = 11.0, gap: float = 8.0, indent: float = 0.0,
             lineheight: float = 1.45) -> None:
        if not content.strip():
            return
        left = _MARGIN + indent
        while True:
            if self.y >= _BOTTOM - size:
                self.new_page()
            rect = self.fitz.Rect(left, self.y, _PAGE_W - _MARGIN, _BOTTOM)
            rc = self.page.insert_textbox(
                rect, content, fontname="china-s", fontsize=size, lineheight=lineheight
            )
            if rc >= 0:
                self.y += (rect.height - rc) + gap
                return
            if self.y <= _MARGIN + 0.5:
                # 单段超过整页：截断兜底，避免死循环
                self.page.insert_textbox(
                    rect, content[:4000], fontname="china-s", fontsize=size, lineheight=lineheight
                )
                self.y = _BOTTOM
                return
            self.new_page()

    def rule(self) -> None:
        if self.y >= _BOTTOM - 12:
            self.new_page()
        self.page.draw_line(
            self.fitz.Point(_MARGIN, self.y), self.fitz.Point(_PAGE_W - _MARGIN, self.y),
            color=(0.7, 0.7, 0.75), width=0.8,
        )
        self.y += 12

    def bytes(self) -> bytes:
        out = self.doc.tobytes()
        self.doc.close()
        return out


def pdf_bytes(md: str) -> bytes:
    p = _Pdf()
    for b in parse_blocks(md):
        kind = b[0]
        if kind == "h":
            size = {1: 20.0, 2: 16.0, 3: 14.0}.get(b[1], 12.5)
            p.text(strip_inline(b[2]), size=size, gap=10.0)
        elif kind == "p":
            p.text(strip_inline(b[1]), size=11.0, gap=9.0)
        elif kind == "ul":
            for it in b[1]:
                p.text("• " + strip_inline(it), size=11.0, gap=4.0, indent=16.0)
            p.y += 4
        elif kind == "ol":
            for n, it in enumerate(b[1], 1):
                p.text(f"{n}. " + strip_inline(it), size=11.0, gap=4.0, indent=16.0)
            p.y += 4
        elif kind == "quote":
            p.text("│ " + strip_inline(b[1]), size=11.0, gap=9.0, indent=14.0)
        elif kind == "code":
            p.text(b[1], size=9.5, gap=9.0, indent=14.0)
        elif kind == "table":
            rows = b[1]
            cols = max(len(r) for r in rows)
            widths = [0] * cols
            for r in rows:
                for i in range(cols):
                    widths[i] = max(widths[i], len(strip_inline(r[i])) if i < len(r) else 0)
            for r in rows:
                cells = [strip_inline(r[i]) if i < len(r) else "" for i in range(cols)]
                line = "  |  ".join(c.ljust(widths[i]) for i, c in enumerate(cells))
                p.text(line, size=10.0, gap=3.0)
            p.y += 6
        elif kind == "hr":
            p.rule()
    return p.bytes()


# --------------------------------------------------------------- 分发
FORMAT_BY_EXT = {
    "pdf": ("pdf", "application/pdf"),
    "docx": ("docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    "md": ("md", "text/markdown; charset=utf-8"),
    "markdown": ("md", "text/markdown; charset=utf-8"),
    "txt": ("txt", "text/plain; charset=utf-8"),
    "html": ("html", "text/html; charset=utf-8"),
}


def render(fmt: str, md: str, title: str = "文档") -> tuple[bytes, str, str]:
    """按目标格式渲染，返回 (字节, MIME, 扩展名)。"""
    if fmt not in FORMAT_BY_EXT:
        raise ValueError(f"不支持的格式：{fmt}")
    ext, mime = FORMAT_BY_EXT[fmt]
    if ext == "pdf":
        return pdf_bytes(md), mime, ext
    if ext == "docx":
        return docx_bytes(md), mime, ext
    if ext == "html":
        return html_bytes(md, title), mime, ext
    if ext == "md":
        return md_bytes(md), mime, ext
    return txt_bytes(md), mime, ext


def download_name(original_name: str | None, doc_type: str, revision: int, ext: str) -> str:
    """下载文件名 = 原文件名（去扩展名）+ (V版本号) + 目标扩展名。"""
    base = Path(original_name or f"文档.{doc_type}").name
    stem = Path(base).stem or "文档"
    return f"{stem}(V{revision}).{ext}"
